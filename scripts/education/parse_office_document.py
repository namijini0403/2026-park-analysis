"""Bounded extraction only. Does not execute macros, formulas or document instructions."""
import base64, csv, io, json, re, struct, sys, zipfile, zlib
from pathlib import Path
from defusedxml import ElementTree as ET

MAX_BYTES=15*1024*1024
MAX_TEXT=500000
MAX_ROWS=15000

def extract(name, data):
    if not data or len(data)>MAX_BYTES: raise ValueError('파일은 15MB 이하이어야 합니다.')
    ext=Path(name).suffix.lower()
    tables, paragraphs, warnings=[],[],[]
    def table(label, rows):
        rows=[[str(v) if v is not None else '' for v in row] for row in rows]
        rows=[r for r in rows if any(v.strip() for v in r)]
        if rows:
            if len(rows)>MAX_ROWS or max(map(len,rows))>150: raise ValueError('표 크기 한도(15,000행·150열) 초과')
            tables.append({'name':label,'rows':rows})
    if ext in ('.xlsx','.docx','.hwpx'):
        z=zipfile.ZipFile(io.BytesIO(data))
        if len(z.infolist())>3000 or sum(i.file_size for i in z.infolist())>60*1024*1024: raise ValueError('압축 문서 크기 한도 초과')
        if ext=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
            for sheet in book:
                if sheet.max_row>MAX_ROWS or sheet.max_column>150: raise ValueError('시트 크기 한도 초과')
                table(sheet.title,sheet.iter_rows(values_only=True))
            book.close()
            warnings.append('수식은 실행하지 않습니다. 수식 셀은 분석에서 결측으로 처리합니다.')
        else:
            files=['word/document.xml'] if ext=='.docx' else sorted(n for n in z.namelist() if re.fullmatch(r'Contents/section\d+\.xml',n))
            if not files: raise ValueError('본문 XML이 없습니다.')
            for file in files:
                root=ET.fromstring(z.read(file))
                local=lambda element:element.tag.rsplit('}',1)[-1]
                text=lambda element:''.join(n.text or '' for n in element.iter() if local(n)=='t')
                for node in root.iter():
                    if local(node)=='p':
                        value=text(node).strip()
                        if value: paragraphs.append(value)
                    if local(node)=='tbl':
                        rows=[]
                        for row in node:
                            if local(row)!='tr':continue
                            rows.append([text(cell) for cell in row if local(cell)=='tc'])
                        table(f'{file} 표 {len(tables)+1}',rows)
            warnings.append('병합 셀·여러 줄 머리글은 미리보기에서 열 위치를 확인하세요.')
    elif ext=='.xls':
        import xlrd
        book=xlrd.open_workbook(file_contents=data,on_demand=True)
        for sheet in book.sheets(): table(sheet.name,(sheet.row_values(i) for i in range(sheet.nrows)))
        book.release_resources()
        warnings.append('XLS 수식은 파일에 저장된 캐시 값으로 읽습니다. 저장 시점을 확인하세요.')
    elif ext=='.hwp':
        import olefile
        with olefile.OleFileIO(io.BytesIO(data)) as doc:
            header=doc.openstream('FileHeader').read()
            if not header.startswith(b'HWP Document File'): raise ValueError('지원하는 HWP 5 형식이 아닙니다.')
            flags=struct.unpack_from('<I',header,36)[0]
            if flags & 6: raise ValueError('암호·배포용 HWP는 해제된 HWP/HWPX 파일로 넣어 주세요.')
            for stream in sorted(p for p in doc.listdir() if p[0]=='BodyText' and p[-1].startswith('Section')):
                raw=doc.openstream(stream).read()
                if flags & 1:
                    decoder=zlib.decompressobj(-15);raw=decoder.decompress(raw,MAX_BYTES+1)
                    if len(raw)>MAX_BYTES or decoder.unconsumed_tail: raise ValueError('HWP 압축해제 크기 한도 초과')
                pos=0
                while pos+4<=len(raw):
                    record=struct.unpack_from('<I',raw,pos)[0];pos+=4
                    tag=record&1023;size=record>>20
                    if size==4095:
                        if pos+4>len(raw):raise ValueError('손상된 HWP')
                        size=struct.unpack_from('<I',raw,pos)[0];pos+=4
                    if pos+size>len(raw):raise ValueError('손상된 HWP 레코드')
                    payload=raw[pos:pos+size];pos+=size
                    if tag==67:
                        # HWP inline controls occupy eight UTF-16 code units.
                        chars=[];i=0
                        while i+2<=len(payload):
                            code=struct.unpack_from('<H',payload,i)[0]
                            if code<32:
                                chars.append(' ')
                                i+=16 if code in (1,2,3,4,5,6,7,8,9,11,12,14,15,16,17,18,19,20,21,22,23) else 2
                            else:chars.append(payload[i:i+2].decode('utf-16le',errors='replace'));i+=2
                        value=''.join(chars).strip()
                        if value:paragraphs.append(value)
        warnings.append('HWP 본문·표 셀의 텍스트 근거를 추출했습니다. 수치 비교용 표 구조는 HWPX 또는 Excel로 저장하여 넣어 주세요.')
    elif ext in ('.csv','.txt'):
        try: text=data.decode('utf-8-sig')
        except UnicodeDecodeError: text=data.decode('cp949')
        if ext=='.csv':table(name,csv.reader(io.StringIO(text)))
        else:paragraphs=text.splitlines()
    elif ext=='.doc':raise ValueError('구형 DOC는 DOCX로 저장한 뒤 넣어 주세요.')
    else:raise ValueError('지원 파일: XLSX, XLS, CSV, DOCX, HWP, HWPX, TXT')
    result={'tables':tables,'paragraphs':paragraphs,'warnings':warnings}
    if len(json.dumps(result,ensure_ascii=False))>MAX_TEXT:raise ValueError('추출 내용이 50만자를 초과합니다. 문서나 시트를 나누어 주세요.')
    if not tables and not paragraphs:raise ValueError('추출 가능한 표·본문이 없습니다. 이미지 문서는 텍스트 문서로 변환해 주세요.')
    return result

if __name__=='__main__':
    sys.stdin.reconfigure(encoding='utf-8');sys.stdout.reconfigure(encoding='utf-8')
    try:
        request=json.load(sys.stdin)
        print(json.dumps(extract(request['name'],base64.b64decode(request['base64'],validate=True)),ensure_ascii=False))
    except Exception as error:
        print(json.dumps({'error':str(error)},ensure_ascii=False));sys.exit(1)
