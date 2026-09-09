"""KEDI school graduation/progression records; never publish pending placeholders."""
import argparse
import hashlib
import json
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd

ROOT=Path(__file__).resolve().parents[2]
DATA=ROOT/'data_processed/education'
CACHE=ROOT/'data/education_sources/kess_school_progression.json'
URL='https://kess.kedi.re.kr/contents/dataset?itemCode=04&menuId=m_02_04_03_02&tabId=m1'
DOWNLOAD_IDS={2025:'2026162291710.xlsx',2026:'2026835615251.xlsx'}
FIELDS={'졸업자_계':'graduates','진학자_계':'advanced','진학률_전체(%)':'published_progression_pct',
        '취업자_계':'employed','입대자_계':'military','기타_계':'other',
        '국내_전문대학_계':'domestic_junior_college','국내_대학_계':'domestic_university',
        '국외_전문대학_계':'overseas_junior_college','국외_대학_계':'overseas_university'}


def extract():
    records=[];sources=[]
    for year in [2025,2026]:
        path=ROOT/f'outputs/education_sources/kess_{year}_schools.xlsx'
        workbook=openpyxl.load_workbook(path,read_only=True,data_only=True)
        sheet=workbook['학교별 주요통계']
        rows=sheet.iter_rows(values_only=True)
        preface=[next(rows) for _ in range(16)]
        header=[str(v or '').replace('\n','').replace(' ','') for v in next(rows)]
        required=['조사기준일','시도','학교급','학교명','학교코드(KEDI)','주소','상태',*FIELDS]
        columns={key:header.index(key) for key in required}
        notes=[str(r[0]) for r in preface if r[0]]
        if not any('고등학교 졸업후 상황' in note and '11월' in note for note in notes):
            raise ValueError('Review changed publication notes before interpreting zeros')
        for row_number,row in enumerate(rows,18):
            if row[columns['시도']]!='인천' or row[columns['학교급']] not in ['중학교','고등학교']:
                continue
            records.append({'year':year,'school_name':row[columns['학교명']], 'school_level':row[columns['학교급']],
                            'kedi_code':str(row[columns['학교코드(KEDI)']]),'address':row[columns['주소']],
                            'school_status':row[columns['상태']], 'reference_date':str(row[columns['조사기준일']]),
                            'source_row':row_number,'metrics':{field:row[columns[key]] for key,field in FIELDS.items()}})
        sources.append({'year':year,'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'notes':notes})
        workbook.close()
    CACHE.write_text(json.dumps({'source_url':URL,'records':records,'sources':sources},ensure_ascii=False,indent=2),encoding='utf-8')


def observation(record):
    result={key:record[key] for key in ['year','reference_date','kedi_code','school_status','source_row']}
    result.update(source_url=URL,metrics=dict(record['metrics']),status='available')
    if record['year']==2026 and record['school_level']=='고등학교':
        result['metrics']={key:value if key=='graduates' else None for key,value in result['metrics'].items()}
        result['status']='progression_pending_publication'
        result['note']='2026년 고등학교 졸업 후 상황은 원본 주석상 11월 이후 갱신 예정. 원본 0을 실적 0으로 해석하지 않음.'
        return result
    values=result['metrics']
    count_keys=['graduates','advanced','employed','military','other']
    if any(not isinstance(values[k],(int,float)) or values[k]<0 for k in count_keys):
        result['status']='missing_or_invalid_observations'
    elif values['graduates']==0:
        values['published_progression_pct']=None
        result['status']='no_graduates'
    else:
        # Published rates rounded to one decimal; do not overwrite disagreements with a derived rate.
        expected=100*values['advanced']/values['graduates']
        rate=values['published_progression_pct']
        if not isinstance(rate,(int,float)) or abs(expected-rate)>.051:
            result['status']='rate_denominator_mismatch'
            result['note']='공표 진학률과 졸업자 대비 진학자 비율 불일치. 원본 수치는 보존하고 해석 보류.'
    return result


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--extract',action='store_true');args=parser.parse_args()
    if args.extract:extract()
    source=json.loads(CACHE.read_text(encoding='utf-8'))
    for entry in source['sources']:
        entry['download_url']=f"https://kess.kedi.re.kr/contents/dataSet/downLoad.do?fileNm={DOWNLOAD_IDS[entry['year']]}&userfileNm=kess_{entry['year']}_schools.xlsx"
    registry=pd.read_csv(DATA/'institutions.csv').to_dict('records')
    by_id={r['학교ID']:r for r in registry}
    aliases=json.loads((DATA/'schoolinfo_verified_aliases.json').read_text(encoding='utf-8'))
    def street(address):return ' '.join(str(address).split('(')[0].split('.')[0].split()[2:])
    local=defaultdict(list)
    for row in registry:local[row['학교명'],row['학교급구분']].append(row)
    schools=defaultdict(list);unmatched=[];seen=set()
    for record in source['records']:
        candidates=local[record['school_name'],record['school_level']]
        basis='인천 내 같은 학교명·학교급 유일 일치'
        if not candidates:
            candidates=[by_id[a['school_id']] for a in aliases.values()
                        if a['disclosure_name']==record['school_name'] and a['school_level']==record['school_level']
                        and str(record['address']).startswith('인천광역시 ')
                        and street(record['address'])==street(a['registry_address'])]
            basis='기존 검증 명칭 별칭 + KEDI 동일 인천 도로명·건물번호 대조'
        if len(candidates)!=1:
            unmatched.append({k:record[k] for k in ['year','school_name','school_level','kedi_code','address']});continue
        sid=candidates[0]['학교ID'];key=(sid,record['year'])
        if key in seen:raise ValueError('Multiple KEDI observations for one school/year; resolve before linking')
        seen.add(key);schools[sid].append({**observation(record),'identity_basis':basis})
    output={'schools':dict(schools),'unmatched':unmatched,'sources':source['sources'],
            'source_sha256':hashlib.sha256(CACHE.read_bytes()).hexdigest(),
            'scope':'교육통계 조사기준일별 학교 졸업 후 상황. 졸업자는 해당 연도 2월 기준. 학교알리미 수시 공시와 조사 기준이 다름.',
            'limitations':['진학률은 진학 등록자/졸업자이며 대학 수준·수능 성적·학교 품질 지표가 아님.',
                           '취업자는 주당 15시간 이상 유급 근무자 기준. 취업률을 임의 산출하지 않음.',
                           '기타는 재수·미확인 등 여러 상황을 포함하며 실패·학업중단으로 해석하지 않음.',
                           '2026 고등학교 졸업 후 상황은 미공개로 처리. 2025년 값을 2026년 값으로 대체하지 않음.']}
    (DATA/'school_progression.json').write_text(json.dumps(output,ensure_ascii=False,separators=(',',':'),allow_nan=False),encoding='utf-8')
    from collections import Counter
    print(json.dumps({'schools':len(schools),'unmatched':len(unmatched),'status':dict(Counter(o['status'] for rows in schools.values() for o in rows))},ensure_ascii=False))


if __name__=='__main__':main()
