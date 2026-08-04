from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[1]
PARENT = ROOT.parent

SOURCE_PPT = PARENT / "incheon_outdoor_equity_v3_6(한희나부장님 수정 5.16).pptx"
OUTPUT_PPT = ROOT / "outputs" / "robust_xai" / "ppt_assets" / "incheon_outdoor_equity_v3_6_robust_xai_inserted.pptx"

SLIDE14_IMAGE = ROOT / "outputs" / "robust_xai" / "ppt_assets" / "slide14_sampling_stability_distribution.png"
SLIDE15_IMAGE = ROOT / "outputs" / "robust_xai" / "ppt_assets" / "slide15_shap_plain_summary.png"


def add_overlay(slide, left, top, width, height, fill=(255, 255, 255)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    shape.line.color.rgb = RGBColor(220, 226, 232)
    shape.line.width = Pt(0.8)
    return shape


def add_title(slide, text, left, top, width, height, size=12):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 45, 61)
    return box


def add_note(slide, text, left, top, width, height, size=7.2):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(88, 101, 116)
    return box


def main():
    if not SOURCE_PPT.exists():
        raise FileNotFoundError(SOURCE_PPT)
    if not SLIDE14_IMAGE.exists() or not SLIDE15_IMAGE.exists():
        raise FileNotFoundError("PPT asset images are missing. Run make_robust_xai_ppt_assets.py first.")

    OUTPUT_PPT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation(SOURCE_PPT)

    # Slide numbers are 1-based in the deck request.
    slide14 = prs.slides[13]
    slide15 = prs.slides[14]

    # Slide 14: add a compact robustness distribution panel in the upper-right area.
    add_overlay(slide14, 8_995_000, 845_000, 2_830_000, 1_925_000)
    add_title(slide14, "1,000회 샘플링 안정성 분포", 9_080_000, 900_000, 2_650_000, 205_000, size=9.5)
    slide14.shapes.add_picture(str(SLIDE14_IMAGE), 9_080_000, 1_120_000, width=2_650_000, height=1_585_000)

    # Slide 15: use the existing Screen 4 slot for the SHAP diagnostic plot.
    add_overlay(slide15, 8_250_000, 2_045_000, 3_190_000, 3_935_000)
    add_title(slide15, "화면 4   SHAP 예측 근거", 8_340_000, 2_135_000, 3_000_000, 260_000, size=11)
    slide15.shapes.add_picture(str(SLIDE15_IMAGE), 8_350_000, 2_480_000, width=2_980_000, height=1_710_000)
    add_note(
        slide15,
        "SHAP은 최종 추천 순위가 아니라 미래 수혜 아동 수 예측값의 변수별 근거를 보여준다.",
        8_370_000,
        4_330_000,
        2_920_000,
        410_000,
    )
    add_note(
        slide15,
        "일반인 설명: 어느 요인이 예측값을 얼마나 밀어 올리거나 낮췄는지 비중(%)으로 확인한다.",
        8_370_000,
        4_730_000,
        2_920_000,
        460_000,
    )

    prs.save(OUTPUT_PPT)
    print(OUTPUT_PPT)


if __name__ == "__main__":
    main()
