"""Build the Word report from saved live HTTP responses and browser captures."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

ROOT=Path(__file__).resolve().parents[2]
DATA=ROOT/'outputs/hitl-release-validation/live'
report=json.loads((DATA/'report.json').read_text(encoding='utf-8'))
audit=json.loads((DATA/'api-audit.json').read_text(encoding='utf-8'))
deployment=json.loads((DATA/'deployment.json').read_text(encoding='utf-8'))
formats=json.loads((DATA/'upload-formats.json').read_text(encoding='utf-8'))
assert formats['passed'] and formats['valid_rows']==12
assert report['passed'] and audit['passed'] and deployment['status']=='SUCCESS'
OUTPUT=ROOT/'contest_plan/HITL_배포와_실제질문_검증보고서_20260912.docx'
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name,size in [('Normal',10.5),('Title',25),('Subtitle',12),('Heading 1',17),('Heading 2',12),('Caption',9)]:
 s=doc.styles[name];s.font.name='Malgun Gothic';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Malgun Gothic')
 s.font.italic=False;s.font.bold=name.startswith('Heading')
 s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=Pt({'Title':33,'Subtitle':17,'Heading 1':22,'Heading 2':16,'Caption':12}.get(name,14.5))
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for grid in list(sec._sectPr.iter(qn('w:docGrid'))):grid.getparent().remove(grid)
doc.core_properties.title='HITL 정책 분석 배포와 실제 질문 검증 보고서'
doc.core_properties.subject='질문별 요소 선택과 설명 가능한 분석 결과의 운영 검증'
doc.core_properties.author='공공데이터 공모전 프로젝트'

def p(text,style=None):return doc.add_paragraph(text,style)
def heading(text):doc.add_heading(text,1)
def page():doc.add_page_break()
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
 t._tbl.tblPr.append(borders)
 for i,(h,w) in enumerate(zip(headers,widths)):t.columns[i].width=Inches(w);t.rows[0].cells[i].text=h
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
 for i,row in enumerate(t.rows):
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E4EDF3' if i==0 else ('F7F9FB' if i%2==0 else 'FFFFFF'));tcpr.append(shd)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    m=OxmlElement('w:'+edge);m.set(qn('w:w'),'85');m.set(qn('w:type'),'dxa');margins.append(m)
   tcpr.append(margins)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3);para.paragraph_format.line_spacing=Pt(12.5)
    for run in para.runs:run.font.size=Pt(9);run.bold=i==0
 p('')
 return t
def picture(file,caption,maxh=3.15):
 file=DATA/file
 with Image.open(file) as im:w,h=im.size
 width=min(6.6,maxh*w/h)
 para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 para.paragraph_format.line_spacing=1.0
 para.paragraph_format.keep_with_next=True
 run=para.add_run();inline=run.add_picture(str(file),width=Inches(width));inline._inline.docPr.set('descr',caption)
 para.paragraph_format.space_after=Pt(3);p(caption,'Caption')

p('HITL 정책 분석 배포와\n실제 질문 검증 보고서','Title')
p('교육 공공데이터 AI 활용대회  |  2026년 9월 12일','Subtitle')
p('질문을 바로 학교 지원 순위로 바꾸지 않고, 사용자가 고려 요소와 중요도를 정한 뒤 원자료·비교집단·확인 조건을 함께 검토하는 흐름을 운영 앱에 적용했다. 42개 자료 영역의 실제 HTTP 요청과 8가지 복합 질문의 브라우저 실행을 확인했다.')
p('등교길 안전인력 질문에는 학생 수·통학구역·공사장 자료를 후보로 제시하지만, 통학 시간대 사고·통행량·기존 인력 자료가 없으므로 어느 학교에 인력을 추가할지 확정하지 않는다. 도서관 질문에는 실제 장서와 거리의 분포 및 선택 학교의 백분위를 제시한다. 결론의 한계를 숨기지 않는 것이 이 분석의 핵심이다.')
table(['검증 항목','확인 결과'],[
 ['운영 주소',report['base']],['배포 상태',deployment['status']],['배포 식별자',deployment['id']],['실행 범위',f"{len(audit['domains'])}개 자료 영역과 {len(report['cases'])}개 복합 질문"],['브라우저','PC 및 모바일 화면, 가중치 변경과 재계산, 자바스크립트 오류 없음'],['원본 대응','배포된 화면 파일 5개 SHA256이 배포 복사본과 일치']],[1.35,5.25])
heading('확인된 기능과 남은 판단')
p('요소 추가·제외, 낮은 값 또는 높은 값의 검토 방향, 중요도 1·2·3, 여러 첨부 요소, 붙여넣은 표와 현장 메모, 분포와 백분위, 근거 저장을 지원한다. 자료의 결측·중복·다른 연도·도서지역을 분리한다.')
p('가중치는 요소의 검토 비중이다. 학교별 합산점수나 투자 순위는 만들지 않는다. 이용 대상·안전·실행·출입구 경로는 가중치로 상쇄할 수 없으며, 현재 HITL 입력만으로 이 조건을 검증 완료로 바꾸는 기능은 제공하지 않는다.')

page();heading('질문에서 분석 결과까지의 구조')
table(['단계','실제 처리'],[
 ['1 질문과 범위','학교명·학교급·연도·지역을 읽고 등록된 자료와 목적별 확인 조건을 제시한다.'],
 ['2 사용자 선택','42개 영역에서 요소를 추가·제외하고 반영 방향과 중요도를 정한다. 처음부터 확정된 요소는 없다.'],
 ['3 추가 자료','CSV·TSV·Excel의 연결 열, JSON 또는 붙여넣은 표를 추가한다. 현장 메모는 확인 조건으로 기록한다.'],
 ['4 서버 검증','등록 요소를 다시 확인하고 결측·중복·연도·단위·학교급과 도서지역을 분리한다.'],
 ['5 계산과 설명','요소별 원값 분포·중앙값·백분위와 학교별 관측표를 표시하고 출처 해시를 기록한다.'],
 ['6 재검토','조건 변경 시 이전 결과 버튼을 비활성화한다. 재계산 후 새 조건과 근거를 저장한다.']],[1.25,5.35])
p('후보 탐색은 보유 자료 목록의 키워드와 명시적인 목적별 규칙을 사용한다. 새로운 언어모델 호출로 지표를 만들어내는 방식은 아니다. 사용자는 추천 후보 외에도 전체 목록에서 요소를 선택할 수 있다.')
picture('library-result.png','운영 앱의 실제 도서관 질문 화면  요소 선택과 근거 패널',3.9)

notes={
'safety':('등교길 안전인력 지원','학생 규모는 참고 관측이다. 공사장 위치와 공원 경로를 실제 통학 위험이나 인력 부족으로 치환하지 않았다.','등하교 시간대 통행량, 사고 및 아차사고, 횡단시설, 기존 인력과 배치 가능시간을 확보해야 한다.'),
'library':('도서관 지원의 여러 근거','장서 총수와 학생당 장서는 서로 다른 관측이다. 학생당 장서를 새 학생 수로 재계산하지 않고 공시값을 유지한다. 외부 도서관 거리는 직선거리로 표시한다.','장서 최신성, 대출·이용량, 사서 근무시간과 외부 도서관의 아동 이용조건을 확인해야 한다.'),
'park':('야외활동과 공원 공동 이용','공원 수, 공유 학교 수와 우회율을 병렬로 보았다. 우회율이 곧 안전 점수라는 해석은 하지 않는다. 중요도 미설정 상태에서는 여러 관점을 함께 표시한다.','실제 출입구, 통행 허용, 시설 안전과 공동 이용 시간대를 확인해야 한다.'),
'care':('지역 돌봄 확대 검토','남동구와 초등학교 범위를 반영했다. 돌봄 시설 기록과 생활권 인구를 학생 수와 합산하지 않으며 정원 부족을 추정하지 않는다.','대기자 수, 정원, 이용 자격, 운영시간과 실제 이동경로를 추가해야 한다.'),
'future':('미래 시설 수요와 불확실성','학교 학생 전망, 학교 예측 오차와 재개발 기록을 구분했다. 예측과 관측은 다른 집단이며 모형 오차를 학교 부족 점수로 바꾸지 않는다.','학구 조정, 사업 추진 일정과 현재 정원 등을 확인하고 여러 시나리오로 검토해야 한다.'),
'curriculum':('고등학교 공동교육과정 참여','고등학교 범위로 전환하고 개설 과목 기록과 학교 연결망을 함께 제시했다. 연결망의 거리는 직선 관계이며 실제 참여나 통학 소요시간이 아니다.','공고의 대상 학년, 참여 자격, 실제 잔여 정원과 시간표·이동 조건을 확인해야 한다.'),
'island':('강화 지역 독서 지원','강화군 조건을 적용하고 도서·농어촌 자료를 일반 학교와 따로 검토했다. 도서관 직선거리가 실제 이동 가능성을 보장하지 않는다.','교통편, 운영시간과 실제 이동 경로를 확인해야 한다. 자료 미확보는 시설 부재를 뜻하지 않는다.'),
'upload':('사용자 현장 자료 추가','브라우저에서 학교명과 시험 통행량 3행을 붙여넣었다. 120·80·100은 업로드 기능 검증용 시험값이며 실제 조사 결과가 아니다. 연결 검증 후 사용자 관측으로만 제시했다.','실제 조사자료는 조사일·시간대·방법·학교 식별자를 포함해야 한다. 메모가 안전 검증 완료로 바뀌지 않는 것을 확인했다.')}
for i,c in enumerate(report['cases'],1):
 item=json.loads((DATA/(c['id']+'.json')).read_text(encoding='utf-8'));result=item['result'];title,meaning,needed=notes[c['id']]
 page();heading(f'사례 {i} {title}')
 p('실제 입력 질문').runs[0].bold=True;p(c['question'])
 choices=[]
 for f in result['review']['factors']:
  direction={'observe':'양쪽 관점','lower':'낮은 값 검토','higher':'높은 값 검토'}[f['direction']]
  weight='미설정' if f['weight'] is None else str(f['weight'])
  share='미설정' if f['share_percent'] is None else str(f['share_percent'])+'%'
  choices.append(f"{f['label']} ({direction}, 중요도 {weight}, 비중 {share})")
 p('사용자가 선택한 요소와 반영 방식  '+' · '.join(choices))
 p('앱의 실제 응답').runs[0].bold=True;p(result['summary'])
 stats=[s for s in result['visual']['sections'] if s.get('statistics')]
 if c['id']=='future':stats=[s for s in stats if '2031' in s['title'] and '예측' in s['title']]
 focused=[s for s in stats if s['statistics'].get('selected')]
 rows=[]
 for s in (focused or stats)[:3]:
  d=s['statistics'];own=d.get('selected');label=s['title'].replace(' · ',' / ')
  rows.append([label,str(d['n']),str(d['median']),str(own['value']) if own else '선택 학교 없음'])
 if rows:table(['지표와 비교 범위','유효 건수','중앙값','선택 학교 값'],rows,[3.4,.8,1.0,1.4])
 chart=c['id']+'-chart.png'
 if (DATA/chart).exists():picture(chart,'운영 앱의 2031년 학생 수 예측 분포' if c['id']=='future' else '사용자 시험자료 3행의 실제 분포' if c['id']=='upload' else '운영 앱에서 캡처한 첫 번째 수치 분포  해당 범위의 유효 관측만 집계',2.65)
 p('해석  '+meaning);p('추가 확인  '+needed)
 p(f"검증 기록  선택 요소 {len(c['factors'])}개 · 근거 구역 {c['sections']}개 · API 및 화면 확인",'Caption')

page();heading('계산 규칙과 검증 결과')
p('중요도 합계가 8이고 한 요소의 중요도가 3이면 그 요소의 검토 비중은 37.5%다. 모든 중요도를 비우면 기본 가중치를 넣지 않고 낮은 값과 높은 값 양쪽 관점을 제공한다. 일부만 중요도를 설정하면 첨부를 포함한 나머지 요소도 지정하도록 안내한다.')
p('값 이하 비율은 해당 값 이하의 유효 관측 수를 분모의 유효 관측 수로 나눈 값이다. 값 이상 비율도 같은 방식이며 동점을 포함한다. 낮은 값을 검토하도록 선택하면 값 이상 비율을, 높은 값을 검토하도록 선택하면 값 이하 비율을 검토 백분위로 사용한다. 두 비율의 합이 100%를 넘을 수 있다.')
lib=json.loads((DATA/'library.json').read_text(encoding='utf-8'))
book=next(s for s in lib['result']['visual']['sections'] if '장서 총수' in s['title'] and s.get('statistics',{}).get('selected'))
b=book['statistics'];own=b['selected']
p(f"도서관 실제 계산 예  {own['name']} 장서 {own['value']:,}권. 이 비교집단의 유효 관측은 {b['n']}건이고 중앙값은 {b['median']:,}권이다. 이 값 이하 비율은 {own['percentile']}%, 이 값 이상 비율은 {own['top_percent']}%다. 전체 등록 학교 전수를 확보했다는 뜻은 아니다.")
table(['검증과 보완','확인한 내용'],[
 ['목적별 후보 보완','등교 안전 질문에 통학구역·공사장·학생 수를 연결하고 공원 경로와 통학 경로의 차이를 명시했다.'],
 ['여러 요소의 연결','학교명 순으로 학교별 요소 관측을 함께 표시한다. 합산점수나 투자 순위는 없다.'],
 ['화면 밀도 보완','다요소 표는 앞 12행, 개별 표는 앞 20행을 표시하고 전체 자료가 JSON에 있음을 안내한다.'],
 ['운영 검증 중 재작업','기록형 요소 막대에 자료 영역명을 표시했다. 첨부 상태 열의 잘못된 인덱스를 수정하고 3행 전부와 백분위를 재검증했다.'],
 ['전체 영역','42개 영역 각각의 요소 목록과 실제 분석 HTTP 응답을 확인했다. 수치가 없는 영역은 기록·지도·출처를 표시한다.'],
 ['정책 회귀 검사','결측·중복 제외, 독립 안전 조건, 기본 점수 없음, 과거 저장값의 근거 없는 추천 복원 방지.'],
 ['사용자 변경','가중치 변경 후 이전 결과 비활성화와 재계산, 첨부 시험자료 연결, PC·모바일 및 오류 여부를 확인했다.'],
 ['배포 검증','특정 배포의 SUCCESS 상태, 운영 파일 해시, 실제 질문 8개와 자료 영역 42개의 응답을 확인했다.']],[1.4,5.2])
page();heading('추가 검증과 재현 방법')
p('재현 스크립트는 tests/verify_hitl_release.cjs와 scripts/deploy/verify_hitl_live.cjs다. 질문·선택·전체 응답·실행 시각·화면은 outputs/hitl-release-validation/live에 보관한다. 배포 복사본에는 파일별 SHA256 목록이 있다. 다른 창의 작업 파일은 덮어쓰지 않고 별도 배포 복사본을 사용했다.')
p('파일 형식 추가 검증 질문  '+formats['question'])
p('CSV·TSV·XLSX 파일과 JSON 붙여넣기 4개를 한 검토안에 추가했다. 12개 시험 행이 모두 연결되고 각 첨부 분포에 반영됐다. 기본 학생 수를 포함한 5개 요소에 명시적 중요도를 적용했다. 재현 스크립트는 tests/verify_hitl_upload_formats.cjs다.')
picture('upload-formats.png','서로 다른 형식의 첨부 4개를 한 검토안에 추가한 운영 화면',3.8)
p('안전인력 배치 효과, 예산 최적 배분, 검증 완료 후보의 정책 우선순위를 이번 결과가 증명하는 것은 아니다. 현재 확보된 자료로 확인 가능한 관측과 추가로 확인해야 할 조건을 제공한다. 현장 자료를 확보한 후에도 최종 판단은 담당자가 수행한다.')
doc.save(OUTPUT)
print(OUTPUT)
